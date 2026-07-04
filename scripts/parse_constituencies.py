#!/usr/bin/env python3
"""Parse INEC 'Electoral Constituencies of the Federation' state XLSX files into
JSON: senatorial districts, federal constituencies and state constituencies, each
with the LGA (senatorial/federal) or RA/ward (state) composition.

Usage: python3 parse_constituencies.py <dir-of-xlsx> <out.json>
Each state file is one workbook with 3 sheets (SENATORIAL / FEDERAL / STATE).
"""
import sys, os, json, re
import openpyxl

CODE_RE = re.compile(r'((?:SD|FC|SC)\s*/\s*\d+\s*/\s*[A-Z]+)', re.I)

def clean(v):
    if v is None: return ""
    return re.sub(r'\s+', ' ', str(v)).strip()

def split_name(raw):
    """'Abia North\\nSD/001/AB' -> ('Abia North', 'SD/001/AB')"""
    raw = clean(raw)
    m = CODE_RE.search(raw)
    code = re.sub(r'\s*', '', m.group(1)).upper() if m else None
    name = CODE_RE.sub('', raw).replace('/', ' ').strip() if m else raw
    name = re.sub(r'\s+', ' ', name).strip(' /')
    return name, code

def grid_from_sheet(ws):
    """Return a dict {(r,c): value} with merged cells forward-filled to every
    cell in the merged range (top-left value)."""
    g = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            g[(r, c)] = clean(ws.cell(r, c).value)
    for rng in ws.merged_cells.ranges:
        val = clean(ws.cell(rng.min_row, rng.min_col).value)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if not g.get((r, c)): g[(r, c)] = val
    return g

def parse_sheet(ws):
    g = grid_from_sheet(ws)
    maxr, maxc = ws.max_row, ws.max_column
    mode = None; name_c = None; comp_c = None
    entries = {"SENATORIAL": [], "FEDERAL": [], "STATE": []}
    cur = None
    for r in range(1, maxr + 1):
        rowvals = [g.get((r, c), "") for c in range(1, maxc + 1)]
        upper = " | ".join(rowvals).upper()
        # detect a header row and its columns
        hn = hc = None; hmode = None
        for c in range(1, maxc + 1):
            U = g.get((r, c), "").upper()
            if "NAME OF SENATORIAL DISTRICT" in U: hmode, hn = "SENATORIAL", c
            elif "NAME OF FEDERAL CONSTITUENCY" in U: hmode, hn = "FEDERAL", c
            elif "NAME OF STATE CONSTITUENCY" in U: hmode, hn = "STATE", c
            if "LGA COMPOSITION" in U or "RA COMPOSITION" in U: hc = c
        if hn and hc:
            mode, name_c, comp_c, cur = hmode, hn, hc, None
            continue
        if not mode or not name_c or not comp_c:
            continue
        nm = g.get((r, name_c), ""); comp = g.get((r, comp_c), "")
        if nm.upper().startswith("TOTAL") or comp.upper().startswith("TOTAL"):
            continue
        if nm and nm.upper() != "S/N":
            name, code = split_name(nm)
            # merged name repeats on continuation rows -> only open a new entry
            # when the name actually changes
            if not cur or cur["_raw"] != nm:
                cur = {"name": name, "code": code, "lgas": [], "_raw": nm}
                entries[mode].append(cur)
        if comp and cur:
            for part in re.split(r'[\n,]+', comp):
                p = part.strip()
                if p and p.upper() != "TOTAL" and p not in cur["lgas"]:
                    cur["lgas"].append(p)
    for k in entries:
        for e in entries[k]: e.pop("_raw", None)
    return entries

def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    out = {"SENATORIAL": [], "FEDERAL": [], "STATE": []}
    for ws in wb.worksheets:
        part = parse_sheet(ws)
        for k in out: out[k].extend(part[k])
    return out

def state_from_filename(fn):
    base = os.path.splitext(os.path.basename(fn))[0]
    base = re.sub(r'^INEC[-_]?', '', base, flags=re.I)
    return re.sub(r'[-_]+', ' ', base).strip().upper()

def main():
    src, out = sys.argv[1], sys.argv[2]
    files = sorted(f for f in os.listdir(src) if f.lower().endswith(('.xlsx', '.xls')))
    result = {}
    for f in files:
        state = state_from_filename(f)
        result[state] = parse_workbook(os.path.join(src, f))
        s = result[state]
        print(f"{state}: senatorial={len(s['SENATORIAL'])} federal={len(s['FEDERAL'])} state={len(s['STATE'])}")
    with open(out, "w") as fh: json.dump(result, fh, indent=2)
    print("wrote", out)

if __name__ == "__main__":
    main()
